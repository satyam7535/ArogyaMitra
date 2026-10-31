"""
ArogyaMitra — Report Styling and Formatting Helpers
Provides document setup, typography, borders, table formatting, callouts, and XML fields.
"""

import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

NAVY_BLUE = RGBColor(0x15, 0x43, 0x60)
SLATE_BLUE = RGBColor(0x1B, 0x4F, 0x72)
DARK_GREY = RGBColor(0x21, 0x21, 0x21)
MUTED_GREY = RGBColor(0x54, 0x6E, 0x7A)

def setup_document_geometry(doc):
    """Sets standard academic margins: 1.25 in left (binding), 1.0 in top/bottom/right."""
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.25)
        section.right_margin = Inches(1.0)
        section.page_width = Inches(8.5)
        section.page_height = Inches(11.0)
        section.different_first_page_header_footer = True
        
        # Header setup
        header = section.header
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hr = hp.add_run("ArogyaMitra: AI-Assisted Primary Healthcare Decision Support System")
        hr.font.name = 'Times New Roman'
        hr.font.size = Pt(8.5)
        hr.font.italic = True
        hr.font.color.rgb = MUTED_GREY
        
        # Footer setup with page numbers
        footer = section.footer
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        fr1 = fp.add_run("Page ")
        fr1.font.name = 'Times New Roman'
        fr1.font.size = Pt(9.5)
        fr1.font.color.rgb = MUTED_GREY
        
        # Add Page field
        fldSimple = OxmlElement('w:fldSimple')
        fldSimple.set(qn('w:instr'), 'PAGE')
        fp._p.append(fldSimple)
        
        fr2 = fp.add_run(" of ")
        fr2.font.name = 'Times New Roman'
        fr2.font.size = Pt(9.5)
        fr2.font.color.rgb = MUTED_GREY
        
        # Add NumPages field
        fldSimple2 = OxmlElement('w:fldSimple')
        fldSimple2.set(qn('w:instr'), 'NUMPAGES')
        fp._p.append(fldSimple2)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('w:top', top), ('w:bottom', bottom), ('w:left', left), ('w:right', right)]:
        node = OxmlElement(m)
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_cell_shading(cell, color_hex):
    shading_xml = f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>'
    cell._tc.get_or_add_tcPr().append(parse_xml(shading_xml))

def set_table_borders(table, color="B0BEC5"):
    tblPr = table._tbl.tblPr
    borders_xml = f'''
    <w:tblBorders {nsdecls("w")}>
        <w:top w:val="single" w:sz="6" w:space="0" w:color="{color}"/>
        <w:bottom w:val="single" w:sz="6" w:space="0" w:color="{color}"/>
        <w:insideH w:val="single" w:sz="4" w:space="0" w:color="{color}"/>
        <w:insideV w:val="none"/>
        <w:left w:val="none"/>
        <w:right w:val="none"/>
    </w:tblBorders>
    '''
    tblPr.append(parse_xml(borders_xml))

def add_p(doc, text="", style='Normal', space_before=0, space_after=6, line_spacing=1.5, align=WD_ALIGN_PARAGRAPH.JUSTIFY, bold=False, italic=False, size=12):
    p = doc.add_paragraph(style=style)
    p.alignment = align
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = line_spacing
    if text:
        run = p.add_run(text)
        run.font.name = 'Times New Roman'
        run.font.size = Pt(size)
        run.bold = bold
        run.italic = italic
        run.font.color.rgb = DARK_GREY
    return p

def add_h1(doc, text):
    doc.add_page_break()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(24)
    p.paragraph_format.space_after = Pt(12)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(16)
    run.bold = True
    run.font.color.rgb = NAVY_BLUE
    return p

def add_h2(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(13.5)
    run.bold = True
    run.font.color.rgb = SLATE_BLUE
    return p

def add_h3(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(12)
    run.bold = True
    run.italic = True
    run.font.color.rgb = DARK_GREY
    return p

def add_caption(doc, text, is_table=False):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(8)
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(10)
    run.bold = True
    run.font.color.rgb = MUTED_GREY
    return p

def add_callout(doc, text, title=""):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_shading(cell, "ECEFF1")
    set_cell_margins(cell, top=140, bottom=140, left=200, right=200)
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(0)
    if title:
        rt = p.add_run(f"[{title}] ")
        rt.font.name = 'Times New Roman'
        rt.font.size = Pt(10.5)
        rt.bold = True
        rt.font.color.rgb = NAVY_BLUE
    r = p.add_run(text)
    r.font.name = 'Times New Roman'
    r.font.size = Pt(10.5)
    r.italic = True
    r.font.color.rgb = DARK_GREY
    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_after = Pt(4)

def insert_image(doc, img_path, caption, width_in=5.8):
    if os.path.exists(img_path):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run()
        run.add_picture(img_path, width=Inches(width_in))
        add_caption(doc, caption)
    else:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(f"[IMAGE NOT FOUND: {os.path.basename(img_path)}]")
        run.bold = True
        run.font.color.rgb = RGBColor(0xC6, 0x28, 0x28)
        add_caption(doc, caption)
