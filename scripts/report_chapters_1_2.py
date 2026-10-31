"""
ArogyaMitra — Report Chapters 1 & 2 Builder
Constructs Chapter 1 (Introduction & Literature Review) and Chapter 2 (System Requirements & Threat Modeling).
"""

from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from report_styles import add_p, add_h1, add_h2, add_h3, add_caption, add_callout, set_cell_margins, set_cell_shading, set_table_borders, NAVY_BLUE, SLATE_BLUE, MUTED_GREY
from report_data import LITERATURE_SUMMARY, COMPARISON_WITH_EXISTING_WORK

def add_chapter_1(doc):
    add_h1(doc, "Chapter 1: Introduction & Literature Review")
    
    add_h2(doc, "1.1 Motivation & Primary Healthcare Challenges in Rural India")
    add_p(doc, "In developing nations with vast geographical dispersion and diverse socio-economic demographics, delivering equitable primary healthcare represents one of the most formidable public health challenges of the twenty-first century. In India, the primary healthcare pyramid rests squarely upon community-level frontline workers, primarily Accredited Social Health Activists (ASHAs) and Auxiliary Nurse Midwives (ANMs) stationed across Health and Wellness Centres (HWCs) and Sub-Centres. These frontline workers operate under severe operational constraints: an acute scarcity of certified medical officers, high patient-to-physician ratios that frequently exceed 1:10,000 in rural districts, rudimentary diagnostic infrastructure, and significant geographic barriers that isolate peripheral villages from tertiary referral hospitals.")
    
    add_p(doc, "When rural patients present with acute symptom constellations, the frontline health worker is tasked with making high-stakes triage decisions: determining whether an individual can be safely managed with symptomatic outpatient care, requires scheduled primary health centre (PHC) physician evaluation, or demands immediate emergency transfer to a district hospital. Errors in this frontline triage process carry severe clinical repercussions. Under-triage of acute conditions—such as atypical myocardial infarction presenting as epigastric discomfort, or acute pediatric dehydration—leads to preventable morbidity and mortality. Conversely, systematic over-triage congests already overburdened secondary and tertiary facilities, imposing catastrophic out-of-pocket travel and consultation expenditures on impoverished rural households.")
    
    add_p(doc, "While digital health initiatives under the Ayushman Bharat Digital Mission (ABDM) have accelerated electronic health record adoption, decision support tools for community health workers remain profoundly deficient. Commercial symptom checkers and diagnostic mobile applications frequently fail in this operating environment. They are typically engineered as monolithic, black-box consumer tools that demand broadband connectivity, assume high digital literacy, fail to provide interpretable clinical rationales, and—critically—lack fail-closed safety architectures. When an underlying algorithmic service experiences network timeout or server failure, black-box systems frequently crash or return silent nulls, leaving frontline workers without guidance at critical junctures.")
    
    add_p(doc, "These systemic challenges establish an urgent national imperative: the engineering of a secure, transparent, and fail-closed AI-assisted clinical decision support system specifically tailored for community healthcare workers. Such a system must not attempt to replace physician diagnosis; rather, it must operate as an auditable cognitive assistant that provides probabilistic differential estimations, enforces deterministic safety overrides for life-threatening emergencies, operates with strict bounded latency, and preserves patient privacy through verifiable defense-in-depth security.")

    add_h2(doc, "1.2 Comprehensive Literature Review")
    add_h3(doc, "1.2.1 Review of Symptom Checkers, Clinical Decision Support Systems, and Calibration in Medicine")
    add_p(doc, "The academic literature evaluating algorithmic decision support and symptom checking tools reveals deep structural contradictions between theoretical predictive accuracy and real-world clinical safety. A cornerstone analysis by Wallace et al. (npj Digital Medicine, 2022) [W] conducted a rigorous systematic review of 10 clinical studies encompassing 48 distinct digital symptom checkers using the standardized Quality Assessment of Diagnostic Accuracy Studies (QUADAS-2) methodology. Their systematic findings provide crucial empirical benchmarks for digital health engineering:")
    
    add_p(doc, "1. Extremely Low Diagnostic Accuracy: Across all evaluated platforms, primary diagnostic accuracy was strikingly poor, ranging between 19.0% and 37.9%. Even when evaluating top-3 or top-5 differential accuracy, symptom checkers consistently failed to match the diagnostic precision of human general practitioners.", space_after=4)
    add_p(doc, "2. Modest Triage Accuracy with Wide Variance: Triage decision accuracy (determining appropriate level of urgency) performed somewhat better, ranging from 48.8% to 90.1%, but exhibited extreme, unpredictable variability across different commercial checkers evaluating identical clinical vignettes.", space_after=4)
    add_p(doc, "3. Swiss-Cheese Model of Safety Hazards: Wallace et al. established that symptom checkers present severe patient safety hazards through an alignment of systemic flaws reminiscent of Reason's Swiss-cheese failure model. Specifically, an incorrect primary diagnosis frequently propagated downstream into dangerous under-triage advice, directing emergency patients toward delayed self-care.", space_after=4)
    add_p(doc, "4. Call for Regulation and Verifiable Oversight: The authors concluded with an unequivocal call for independent clinical validation, transparent algorithmic auditing, and strict medical device regulation, noting that the vast majority of commercial tools operate without published validation data.", space_after=10)

    add_p(doc, "Complementing clinical triage evaluation, the theoretical foundations of probabilistic classifier calibration were fundamentally established by Niculescu-Mizil & Caruana (ICML, 2005) [N]. In their landmark empirical investigation across 10 supervised learning families, the authors demonstrated that classification algorithms optimize class separation boundaries rather than true posterior probabilities. Specifically, they proved that:")
    
    add_p(doc, "1. Distortion in Naive Bayes Estimators: Naive Bayes models push predicted posterior probabilities asymptotically toward the extremes of 0 and 1. This occurs because the conditional feature independence assumption multiplies marginal likelihoods, causing redundant or correlated evidence to artificially inflate the posterior log-odds.", space_after=4)
    add_p(doc, "2. Data Scarcity and Calibration Overfitting: Post-hoc calibration methods—such as Platt scaling (sigmoid fitting) and Isotonic regression—require substantial, independent validation cohorts. On small or sparse medical datasets, Isotonic regression overfits aggressively by memorizing rank orderings, while Platt scaling distorts well-separated likelihoods.", space_after=10)

    add_callout(doc, "Theoretical Insight from Niculescu-Mizil & Caruana (2005): Naive Bayes models produce well-ranked decision margins under zero-one loss (Domingos & Pazzani, 1997) but generate uncalibrated posterior probabilities that cannot be interpreted as true clinical disease likelihoods without substantial independent calibration data.", "Probabilistic Calibration Note")

    add_p(doc, "Table 1.1 synthesizes 10 seminal contributions across clinical symptom checking, Bayesian theory, machine learning interpretability, natural language generation safety, and cryptographic security that directly informed the architecture of ArogyaMitra.")

    # Table 1.1
    t11 = doc.add_table(rows=len(LITERATURE_SUMMARY)+1, cols=4)
    t11.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t11)
    col_widths = [Inches(1.5), Inches(1.1), Inches(2.3), Inches(2.1)]
    
    headers = ["Author & Year", "Venue", "Methodological Summary", "Relevance to ArogyaMitra"]
    for j, h in enumerate(headers):
        cell = t11.cell(0, j)
        cell.width = col_widths[j]
        set_cell_shading(cell, "154360")
        set_cell_margins(cell, top=120, bottom=120, left=100, right=100)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(h)
        r.font.name = 'Times New Roman'
        r.font.size = Pt(9.5)
        r.bold = True
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        
    for i, row in enumerate(LITERATURE_SUMMARY):
        row_cells = t11.rows[i+1].cells
        for j, w in enumerate(col_widths):
            row_cells[j].width = w
            set_cell_margins(row_cells[j], top=80, bottom=80, left=90, right=90)
            if i % 2 == 1:
                set_cell_shading(row_cells[j], "F4F6F7")
        
        # Col 0: Author
        p0 = row_cells[0].paragraphs[0]
        r = p0.add_run(row["author"])
        r.font.name = 'Times New Roman'; r.font.size = Pt(8.5); r.bold = True
        
        # Col 1: Venue
        p1 = row_cells[1].paragraphs[0]
        r = p1.add_run(row["venue"])
        r.font.name = 'Times New Roman'; r.font.size = Pt(8.5); r.italic = True
        
        # Col 2: Method
        p2 = row_cells[2].paragraphs[0]
        r = p2.add_run(row["method"])
        r.font.name = 'Times New Roman'; r.font.size = Pt(8.5)
        
        # Col 3: Relevance
        p3 = row_cells[3].paragraphs[0]
        r = p3.add_run(row["relevance"])
        r.font.name = 'Times New Roman'; r.font.size = Pt(8.5)

    add_caption(doc, "Table 1.1: Comprehensive Literature Review Summary Matrix", is_table=True)

    add_h3(doc, "1.2.2 Comparison with Existing Work")
    add_p(doc, "To highlight the engineering novelties and uncompromising safety boundaries of ArogyaMitra, Table 1.2 contrasts specific clinical concerns documented in literature against conventional symptom checkers and the verifiable implementations within this repository.")

    # Table 1.2
    t12 = doc.add_table(rows=len(COMPARISON_WITH_EXISTING_WORK)+1, cols=4)
    t12.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t12)
    cw12 = [Inches(1.8), Inches(1.3), Inches(2.2), Inches(1.7)]
    
    h12 = ["Clinical / Technical Concern", "Literature Citation", "ArogyaMitra Implementation", "Unvalidated Boundary"]
    for j, h in enumerate(h12):
        cell = t12.cell(0, j)
        cell.width = cw12[j]
        set_cell_shading(cell, "1B4F72")
        set_cell_margins(cell, top=120, bottom=120, left=100, right=100)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(h)
        r.font.name = 'Times New Roman'; r.font.size = Pt(9.5); r.bold = True
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        
    for i, row in enumerate(COMPARISON_WITH_EXISTING_WORK):
        row_cells = t12.rows[i+1].cells
        for j, w in enumerate(cw12):
            row_cells[j].width = w
            set_cell_margins(row_cells[j], top=80, bottom=80, left=90, right=90)
            if i % 2 == 1:
                set_cell_shading(row_cells[j], "F4F6F7")
        
        p0 = row_cells[0].paragraphs[0]
        r = p0.add_run(row["concern"])
        r.font.name = 'Times New Roman'; r.font.size = Pt(8.5); r.bold = True
        
        p1 = row_cells[1].paragraphs[0]
        r = p1.add_run(row["source"])
        r.font.name = 'Times New Roman'; r.font.size = Pt(8.5); r.italic = True
        
        p2 = row_cells[2].paragraphs[0]
        r = p2.add_run(row["response"])
        r.font.name = 'Times New Roman'; r.font.size = Pt(8.5)
        
        p3 = row_cells[3].paragraphs[0]
        r = p3.add_run(row["unvalidated"])
        r.font.name = 'Times New Roman'; r.font.size = Pt(8.5); r.font.color.rgb = RGBColor(0xB7, 0x1C, 0x1C)

    add_caption(doc, "Table 1.2: Systematic Comparison: Clinical Concerns vs ArogyaMitra Implementations", is_table=True)

    add_h2(doc, "1.3 Problem Statement & Design Objectives")
    add_p(doc, "Problem Statement: Community health workers in resource-constrained primary healthcare settings require an auditable, privacy-preserving, and fail-closed decision support tool to triage acute clinical presentations without relying on black-box predictions, ungrounded generative hallucinations, or vulnerable cloud infrastructure.")
    
    add_p(doc, "To solve this problem, ArogyaMitra was engineered according to five inviolable design objectives:", bold=True)
    add_p(doc, "1. Objective 1 — Probabilistic Explainability: Deploy an inference engine where every differential disease score is mathematically traceable to specific symptom attributions, eliminating black-box opacity.", space_after=4)
    add_p(doc, "2. Objective 2 — Deterministic Red-Flag Triage Precedence: Guarantee that clinical red-flag symptoms (e.g., severe chest pain radiating to the left arm, acute dyspnea, altered sensorium) instantly trigger emergency escalation, completely superseding statistical model predictions.", space_after=4)
    add_p(doc, "3. Objective 3 — Guarded LLM Narrative Generation: Confine Large Language Models strictly to formatting human-readable summaries and explanations, enforcing strict per-attempt (10s) and cascade (20s) execution caps backed by deterministic medical templates.", space_after=4)
    add_p(doc, "4. Objective 4 — Fail-Closed System Behavior: In the event of microservice failure, network disruption, or database unreachability, the system must never output misleading fallback predictions; it must reject unsafe operations and return explicit HTTP 503 error states.", space_after=4)
    add_p(doc, "5. Objective 5 — Defense-in-Depth Data Security: Enforce cryptographic protection across every trust boundary, including field-level encryption for patient identifiers, memory-hard credential hashing, and hardened Docker container network isolation.", space_after=12)

    add_h2(doc, "1.4 Faculty Guidance & Architectural Evolution")
    add_p(doc, "The development of ArogyaMitra was profoundly shaped by iterative technical reviews and rigorous critique from project supervisor Dr. Ashish Singh. The initial prototype designed during early development phases was structured as a standard full-stack machine learning demonstrator: a monolithic script exposing a tree-based classifier, an unconstrained generative LLM prompt, and default web application security. Through systematic faculty reviews, several critical architectural vulnerabilities were identified and systematically remediated:")
    
    add_p(doc, "1. Rejection of Spurious Model Calibration Claims: When initial drafts reported high confidence numbers and attempted post-hoc Platt scaling, faculty review directed an in-depth reading of Niculescu-Mizil & Caruana (2005) [N]. Empirical verification revealed that fitting a Platt sigmoid on 304 unique rows severely degraded log loss from 0.0061 to 0.1643. Under faculty direction, all misleading claims of 'calibrated probabilities' were excised. The UI was updated to explicitly inform health workers that displayed scores represent 'uncalibrated posterior estimates from Bernoulli Naive Bayes'.", space_after=6)
    
    add_p(doc, "2. Mandating Fail-Closed Safety Boundaries: Early prototype versions utilized client-side fallback mocks that returned simulated disease predictions when backend services were offline. Faculty review condemned this pattern as clinically catastrophic, pointing to Wallace et al. (2022) [W]. The system was completely re-architected to be fail-closed: eliminating all mock fallbacks and guaranteeing that an ML service outage returns an explicit HTTP 503 MODEL_UNAVAILABLE error with clear instructions to seek emergency care.", space_after=6)
    
    add_p(doc, "3. Replacement of Sampling Explanations with Exact Log-Odds: Initial plans to incorporate LIME (Ribeiro et al., 2016) were rejected due to sampling stochasticity and high latency overhead (exceeding 800 ms per query). Faculty guidance prompted the mathematical derivation of exact closed-form Bayesian log-odds attributions directly from the Bernoulli parameter matrix, achieving sub-millisecond explanation generation to machine precision.", space_after=6)
    
    add_p(doc, "4. Container Decoupling and Secret Hardening: Static analysis detected fallback encryption keys embedded in source files. Under faculty supervision, the architecture was refactored into four isolated Docker microservices connected over a private bridge network, with strict environment variable enforcement and an automated production-guard test verifying that missing encryption keys halt execution immediately.", space_after=12)

    add_h2(doc, "1.5 Report Outline & Organization")
    add_p(doc, "The remainder of this report is structured as follows: Chapter 2 specifies the functional, non-functional, and STRIDE threat modeling requirements governing clinical decision support. Chapter 3 provides an exhaustive technical breakdown of the microservice architecture, mathematical formulation of Bernoulli Naive Bayes, exact log-odds derivation, red-flag screening rules, guarded LLM timeout cascades, cryptographic implementations, and concurrency controls. Chapter 4 presents extensive empirical evaluations, including nested repeated cross-validation across symptom dropout levels, multi-tier load testing at scale, automated test suite verification, and literature-linked discussion. Chapter 5 concludes the report and delineates future research directions. Detailed engineering challenges, CLI reproduction commands, and test suite inventories are provided in Appendices A, B, and C.")

def add_chapter_2(doc):
    add_h1(doc, "Chapter 2: System Requirements & Threat Modeling")
    
    add_h2(doc, "2.1 Functional Requirements")
    add_p(doc, "The ArogyaMitra system is designed to govern the end-to-end clinical triage workflow across frontline community health workers, consulting physicians, and administrative auditors. Table 2.1 specifies the ten core functional requirements (FR-1 through FR-10) implemented across the microservice stack.")

    # Table 2.1
    fr_data = [
        ("FR-1", "Consultation Intake & Demographics", "The system shall capture patient basic demographics (age, gender, anonymized district) and presenting chief complaints, validating inputs against strict Pydantic schemas."),
        ("FR-2", "Controlled Symptom Vocabulary Mapping", "The system shall allow health workers to select symptoms exclusively from a validated vocabulary of 131 clinical features, converting selections into binary indicator vectors."),
        ("FR-3", "Deterministic Red-Flag Screening", "Prior to ML evaluation, the system shall execute deterministic rule matching across emergency conditions; any red-flag match shall immediately assign Emergency triage."),
        ("FR-4", "Bayesian Differential Classification", "The ML microservice shall compute uncalibrated posterior log-odds across 41 disease classes and return the top-3 ranked differential diagnoses within 50 milliseconds."),
        ("FR-5", "Additive Feature Attribution", "The ML engine shall calculate exact positive and negative log-likelihood contributions for each presenting symptom to explain why a disease was ranked highest."),
        ("FR-6", "Guarded LLM Clinical Narrative", "The LangGraph service shall generate patient-friendly clinical summaries and care recommendations using Gemini models, subject to strict confidence-word filtering."),
        ("FR-7", "Deterministic Template Fallback", "If LLM calls timeout (10s per-attempt, 20s total cap), encounter 429 quota exhaustion, or fail guardrail checks, the system shall seamlessly render pre-compiled clinical templates."),
        ("FR-8", "Dual Health Worker Decision Gate", "The UI shall provide frontline workers with an explicit Accept / Escalate decision interface; accepting logs worker sign-off, while escalating routes to the physician queue."),
        ("FR-9", "Server-Side Consultation Pagination", "The backend shall expose indexed, paginated historical consultation endpoints (`GET /consultation/my`) supporting limit and page parameters with lean projection."),
        ("FR-10", "DPDP-Compliant Data Consent & Purge", "The system shall support digital consent recording, data anonymization, and verified record revocation in alignment with the Indian Digital Personal Data Protection (DPDP) Act.")
    ]
    
    t21 = doc.add_table(rows=len(fr_data)+1, cols=3)
    t21.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t21)
    cw21 = [Inches(1.0), Inches(2.2), Inches(3.8)]
    
    h21 = ["Req ID", "Requirement Title", "Technical Specification & Scope"]
    for j, h in enumerate(h21):
        cell = t21.cell(0, j)
        cell.width = cw21[j]
        set_cell_shading(cell, "154360")
        set_cell_margins(cell, top=100, bottom=100, left=90, right=90)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(h)
        r.font.name = 'Times New Roman'; r.font.size = Pt(9.5); r.bold = True
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        
    for i, (rid, rtitle, rdesc) in enumerate(fr_data):
        row_cells = t21.rows[i+1].cells
        for j, w in enumerate(cw21):
            row_cells[j].width = w
            set_cell_margins(row_cells[j], top=70, bottom=70, left=80, right=80)
            if i % 2 == 1:
                set_cell_shading(row_cells[j], "F4F6F7")
        
        p0 = row_cells[0].paragraphs[0]
        r = p0.add_run(rid); r.font.name = 'Times New Roman'; r.font.size = Pt(9); r.bold = True
        p1 = row_cells[1].paragraphs[0]
        r = p1.add_run(rtitle); r.font.name = 'Times New Roman'; r.font.size = Pt(9); r.bold = True
        p2 = row_cells[2].paragraphs[0]
        r = p2.add_run(rdesc); r.font.name = 'Times New Roman'; r.font.size = Pt(8.5)

    add_caption(doc, "Table 2.1: System Functional Requirements Specification (FR-1 to FR-10)", is_table=True)

    add_h2(doc, "2.2 Non-Functional Requirements")
    add_p(doc, "To ensure operational viability in low-bandwidth rural environments, Table 2.2 defines eight non-functional requirements governing latency, reliability, security, and concurrency.")

    # Table 2.2
    nfr_data = [
        ("NFR-1", "Bounded Gateway Latency", "End-to-end API gateway triage latency shall not exceed 1,500 ms at p95 under standard operating load, and shall degrade gracefully under network contention."),
        ("NFR-2", "ML Inference Micro-Latency", "Stateless Bayesian differential inference shall execute in under 10 ms on standard commodity CPU cores without requiring hardware GPU accelerators."),
        ("NFR-3", "Fail-Closed High Availability", "Microservice unavailability shall result in deterministic HTTP 503 error codes within 5,000 ms, never leaking unhandled stack traces or hanging connections."),
        ("NFR-4", "Cryptographic Data Protection", "All patient identifiers and sensitive medical fields in MongoDB shall be encrypted at rest using AES-256-GCM with unique 96-bit initialization vectors per record."),
        ("NFR-5", "Resilient Authentication Hashing", "User authentication credentials shall be protected using Argon2id with 3 iterations, 64 MB memory cost, and 4 degrees of parallelism to prevent offline cracking."),
        ("NFR-6", "Concurrency & State Consistency", "The system shall maintain zero state divergence and prevent lost updates during simultaneous triage responses using Mongoose versioning (`__v`) and atomic updates."),
        ("NFR-7", "Strict Transport Security", "All external network interactions shall enforce HTTPS over TLS 1.3 with automated HSTS, X-Content-Type-Options, and Content Security Policy headers."),
        ("NFR-8", "Containerized Environment Isolation", "Services shall execute as unprivileged processes inside lightweight Docker containers connected via private network bridges with no public database exposure.")
    ]
    
    t22 = doc.add_table(rows=len(nfr_data)+1, cols=3)
    t22.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t22)
    cw22 = [Inches(1.0), Inches(2.2), Inches(3.8)]
    
    h22 = ["Req ID", "Performance Dimension", "Engineering Target & Enforcement"]
    for j, h in enumerate(h22):
        cell = t22.cell(0, j)
        cell.width = cw22[j]
        set_cell_shading(cell, "1B4F72")
        set_cell_margins(cell, top=100, bottom=100, left=90, right=90)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(h)
        r.font.name = 'Times New Roman'; r.font.size = Pt(9.5); r.bold = True
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        
    for i, (rid, rtitle, rdesc) in enumerate(nfr_data):
        row_cells = t22.rows[i+1].cells
        for j, w in enumerate(cw22):
            row_cells[j].width = w
            set_cell_margins(row_cells[j], top=70, bottom=70, left=80, right=80)
            if i % 2 == 1:
                set_cell_shading(row_cells[j], "F4F6F7")
        
        p0 = row_cells[0].paragraphs[0]
        r = p0.add_run(rid); r.font.name = 'Times New Roman'; r.font.size = Pt(9); r.bold = True
        p1 = row_cells[1].paragraphs[0]
        r = p1.add_run(rtitle); r.font.name = 'Times New Roman'; r.font.size = Pt(9); r.bold = True
        p2 = row_cells[2].paragraphs[0]
        r = p2.add_run(rdesc); r.font.name = 'Times New Roman'; r.font.size = Pt(8.5)

    add_caption(doc, "Table 2.2: System Non-Functional Requirements Specification (NFR-1 to NFR-8)", is_table=True)

    add_h2(doc, "2.3 Threat Model & Security Analysis (STRIDE Matrix)")
    add_p(doc, "Healthcare software platforms operating across decentralized network endpoints represent high-value targets for malicious exploitation, eavesdropping, and data tampering. To systematically identify threat vectors and engineer robust countermeasures, ArogyaMitra was subjected to an exhaustive STRIDE threat analysis (Spoofing, Tampering, Repudiation, Information Disclosure, Denial of Service, and Elevation of Privilege). Table 2.3 summarizes the identified threat scenarios, potential clinical impacts, and the specific architectural mitigations verified in the codebase.")

    # Table 2.3
    stride_data = [
        ("Spoofing", "Adversary intercepts network traffic or forges worker credentials to submit fabricated patient encounters.", "High (Clinical Distortion)", "Enforced short-lived signed JWTs (15-min expiry) with HMAC-SHA256, Argon2id password hashing, and strict role validation."),
        ("Tampering", "Adversary modifies triage urgency scores or alters patient symptom vectors in transit.", "Critical (Under-triage)", "Enforced TLS 1.3 encryption across external gateways; internal Docker network isolation; immutable SHA-256 hash chains on triage logs."),
        ("Repudiation", "Health worker denies recording an improper triage decision or clinician denies failing to review an escalation.", "High (Legal Liability)", "Enforced immutable audit logging in MongoDB recording exact ISO timestamps, authenticated user IDs, and client IP addresses."),
        ("Information Disclosure", "Eavesdropper accesses patient PII or unencrypted diagnosis records through database dumps or packet sniffing.", "Critical (Privacy Violation)", "Field-level encryption using AES-256-GCM for patient names and contacts; PII scrubbing before sending text to external LLM APIs."),
        ("Denial of Service", "Adversary floods triage endpoints with HTTP requests, exhausting CPU or Gemini API quotas.", "High (Service Outage)", "Multi-tier rate limiting (100 req/15 min on public endpoints); hard 20s LLM timeout cap with instant fallback to local cached templates."),
        ("Elevation of Privilege", "Frontline health worker accesses administrative dashboards or modifies physician clinical notes.", "High (Unauthorized Action)", "Role-Based Access Control (RBAC) middleware verifying user roles ('worker', 'doctor', 'admin') on every API route.")
    ]
    
    t23 = doc.add_table(rows=len(stride_data)+1, cols=4)
    t23.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t23)
    cw23 = [Inches(1.2), Inches(2.2), Inches(1.3), Inches(2.3)]
    
    h23 = ["STRIDE Category", "Identified Threat Scenario", "Clinical Impact", "Architectural Mitigation"]
    for j, h in enumerate(h23):
        cell = t23.cell(0, j)
        cell.width = cw23[j]
        set_cell_shading(cell, "154360")
        set_cell_margins(cell, top=100, bottom=100, left=90, right=90)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(h)
        r.font.name = 'Times New Roman'; r.font.size = Pt(9.5); r.bold = True
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        
    for i, (scat, sthreat, simpact, smit) in enumerate(stride_data):
        row_cells = t23.rows[i+1].cells
        for j, w in enumerate(cw23):
            row_cells[j].width = w
            set_cell_margins(row_cells[j], top=70, bottom=70, left=80, right=80)
            if i % 2 == 1:
                set_cell_shading(row_cells[j], "F4F6F7")
        
        p0 = row_cells[0].paragraphs[0]
        r = p0.add_run(scat); r.font.name = 'Times New Roman'; r.font.size = Pt(9); r.bold = True
        p1 = row_cells[1].paragraphs[0]
        r = p1.add_run(sthreat); r.font.name = 'Times New Roman'; r.font.size = Pt(8.5)
        p2 = row_cells[2].paragraphs[0]
        r = p2.add_run(simpact); r.font.name = 'Times New Roman'; r.font.size = Pt(8.5); r.bold = True
        if "Critical" in simpact:
            r.font.color.rgb = RGBColor(0xB7, 0x1C, 0x1C)
        elif "High" in simpact:
            r.font.color.rgb = RGBColor(0xE6, 0x51, 0x00)
        p3 = row_cells[3].paragraphs[0]
        r = p3.add_run(smit); r.font.name = 'Times New Roman'; r.font.size = Pt(8.5)

    add_caption(doc, "Table 2.3: STRIDE Threat Model & Architectural Security Controls", is_table=True)

    add_h2(doc, "2.4 Clinical Safety & Ethics Requirements")
    add_p(doc, "The deployment of artificial intelligence in healthcare demands ethical and safety guardrails far more stringent than standard enterprise software. In accordance with guidelines from the World Health Organization (WHO Guidance on Ethics and Governance of AI for Health) and the Indian Council of Medical Research (ICMR Ethical Guidelines for AI in Biomedical Research), ArogyaMitra enforces four inviolable clinical safety principles:")
    
    add_p(doc, "1. Strict Human-in-the-Loop Governance: ArogyaMitra is architected solely as a decision support aid, not an autonomous diagnostic agent. The system cannot initiate clinical treatments, prescribe scheduled pharmaceutical drugs, or finalize patient discharge without explicit review and confirmation by a certified human health worker or medical officer.", space_after=6)
    
    add_p(doc, "2. Absolute Red-Flag Precedence: Statistical machine learning models can fail unpredictably when presented with complex, atypical, or out-of-distribution symptom vectors. Therefore, rule-based emergency detection operates with absolute precedence over statistical predictions, ensuring that life-threatening symptoms trigger immediate emergency transfer regardless of model confidence.", space_after=6)
    
    add_p(doc, "3. Complete Disclosure of Algorithmic Boundaries: To prevent automation bias—wherein health workers place unwarranted trust in algorithmic outputs—the user interface explicitly displays model boundaries, labels scores as uncalibrated estimates, and provides transparent feature contribution breakdowns.", space_after=6)
    
    add_p(doc, "4. Protection Against Synthetic Data Artifacts: The machine learning model was trained on synthetic clinical datasets. While synthetic data eliminates patient privacy risks during development, it lacks the biological noise, comorbidities, and demographic nuances of live clinical populations. The system explicitly incorporates these operational limits into its documentation and consent workflows.", space_after=12)
