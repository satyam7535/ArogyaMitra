"""
ArogyaMitra — Report Chapter 5, References & Appendices Builder (Item 1, 6, 14, 18, 19 Compliant)
Builds Chapter 5 (Conclusion & Future Scope), Formal References [1]-[15],
Appendix A (Engineering Challenges & Fixes), and Appendix B (Reproduction Playbook).
Appendix C is deleted and replaced by Table 4.5 per Item 18.
"""

from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from report_styles import (
    add_p, add_h1, add_h2, add_table_caption,
    set_cell_margins, set_cell_shading, set_table_borders,
    NAVY_BLUE, MUTED_GREY, DARK_GREY
)
from report_data import APPENDIX_A_CHALLENGES

def add_chapter_5(doc):
    add_h1(doc, "Chapter 5: Conclusion & Future Scope")
    
    add_h2(doc, "5.1 Project Summary & Achievements")
    add_p(doc, "In this project, I engineered and evaluated ArogyaMitra, an explainable and fail-closed AI-assisted clinical decision support system prototype for community healthcare workers. Addressing vulnerabilities documented in clinical literature—specifically diagnostic inaccuracy in symptom checkers (Wallace et al. [1]) and uncalibrated posteriors in Naive Bayes (Niculescu-Mizil & Caruana [2])—I achieved the following engineering outcomes:")
    
    add_p(doc, "1. Deduplicated Baseline Benchmark: Identified and removed a 93.8% synthetic duplication artifact in raw public clinical data, establishing a clean 304-row benchmark across 41 disease classes and 131 symptoms.", space_after=3)
    add_p(doc, "2. Robust Bayesian Inference: Evaluated stateless Bernoulli Naive Bayes against tuned gradient-boosted trees under symptom dropout, demonstrating 0.919 accuracy vs. 0.683 for LightGBM at 50% missing data, with an inference execution time of 1.8 ms in internal container benchmarks.", space_after=3)
    add_p(doc, "3. Closed-Form Feature Attribution: Derived exact mathematical log-odds symptom attributions satisfying additive efficiency without sampling approximations.", space_after=3)
    add_p(doc, "4. Fail-Closed Safety Architecture: Decoupled emergency detection from classification via deterministic red-flag overrides; isolated the LLM narrative layer behind 10s/20s timeout caps and confidence filters backed by pre-compiled templates; and replaced client-side mocks with explicit HTTP 503 error states.", space_after=3)
    add_p(doc, "5. Layered Container Security & Verification: Enforced TLS 1.2/1.3, Argon2id hashing, AES-256-GCM field encryption, and optimistic concurrency controls (__v). Achieved zero fixable Critical/High (Trivy --ignore-unfixed) container vulnerabilities and 100% pass rates across 322 automated prototype tests.", space_after=8)

    add_h2(doc, "5.2 Lessons Learned & Engineering Insights")
    add_p(doc, "The iterative design and testing of ArogyaMitra yielded key engineering insights:")
    add_p(doc, "1. Danger of Silent Fallback Mocks: Returning canned predictions during a service outage creates clinical hazards. Designing for fail-closed behavior—explicitly halting unsafe workflows and returning HTTP 503—is essential in healthcare software.", space_after=3)
    add_p(doc, "2. Calibration Limits on Small Cohorts: Demonstrating that Platt scaling degraded log loss on 304 unique rows confirmed that parametric calibration cannot be applied uncritically to small clinical datasets. Scores must be transparently communicated as uncalibrated estimates.", space_after=3)
    add_p(doc, "3. Combined Profiling & Optimization: Profiling revealed that unpaginated queries dominated list latency. Implementing server-side pagination and projection yielded a 60.5% p50 latency reduction, confirming the value of targeted optimization.", space_after=8)

    add_h2(doc, "5.3 Future Work & Roadmap")
    add_p(doc, "Key milestones for future development include:")
    add_p(doc, "1. Prospective Clinical Pilot: Conduct observational studies in Health and Wellness Centres to evaluate concordance between health worker triage recommendations and Medical Officer diagnoses.", space_after=3)
    add_p(doc, "2. ABHA & ABDM Ecosystem Integration: Implement standardized HL7 FHIR interfaces for bidirectional health record synchronization with the Ayushman Bharat Digital Mission.", space_after=3)
    add_p(doc, "3. Multi-Modal Intake: Ingest continuous vital signs (SpO2, pulse rate, automated blood pressure) alongside symptom indicator vectors.", space_after=3)
    add_p(doc, "4. Non-Parametric Calibration on Clinical Cohorts: When genuine multi-thousand patient cohorts become available, train non-parametric calibration models (such as beta calibration) to deliver calibrated risk intervals.", space_after=8)

def add_references(doc):
    add_h1(doc, "References")
    
    refs = [
        ("Wallace, W., Chan, C., Chidambaram, S., et al. (2022). The diagnostic and triage accuracy of digital and online symptom checker tools: a systematic review. npj Digital Medicine, 5:118. https://doi.org/10.1038/s41746-022-00667-w"),
        ("Niculescu-Mizil, A., & Caruana, R. (2005). Predicting good probabilities with supervised learning. In Proceedings of the 22nd International Conference on Machine Learning (ICML '05), pp. 625–632. ACM. https://doi.org/10.1145/1102351.1102430"),
        ("Domingos, P., & Pazzani, M. (1997). On the optimality of the simple Bayesian classifier under zero-one loss. Machine Learning, 29(2), 103–130. https://doi.org/10.1023/A:1007413511361"),
        ("Ke, G., Meng, Q., Finley, T., et al. (2017). LightGBM: A highly efficient gradient boosting decision tree. Advances in Neural Information Processing Systems (NeurIPS 2017), 30, 3146–3154."),
        ("Lundberg, S. M., & Lee, S. I. (2017). A unified approach to interpreting model predictions. Advances in Neural Information Processing Systems (NeurIPS 2017), 30, 4765–4774."),
        ("Ji, Z., Lee, N., Frieske, R., et al. (2023). Survey of hallucination in natural language generation. ACM Computing Surveys, 55(12), 1–38. https://doi.org/10.1145/3571730"),
        ("Singhal, K., Azizi, S., Tu, T., et al. (2023). Large language models encode clinical knowledge. Nature, 620(7972), 172–180. https://doi.org/10.1038/s41586-023-06291-2"),
        ("Biryukov, A., Dinu, D., & Khovratovich, D. (2016). Argon2: New generation of memory-hard functions for password hashing. In IEEE European Symposium on Security and Privacy (EuroS&P), pp. 292–302. IEEE."),
        ("Kung, H. T., & Robinson, J. T. (1981). On optimistic methods for concurrency control. ACM Transactions on Database Systems (TODS), 6(2), 213–226. https://doi.org/10.1145/319566.319567"),
        ("Ribeiro, M. T., Singh, S., & Guestrin, C. (2016). \"Why should I trust you?\": Explaining the predictions of any classifier. In Proceedings of the 22nd ACM SIGKDD (KDD '16), pp. 1135–1144. ACM."),
        ("Ministry of Law and Justice, Government of India. (2023). The Digital Personal Data Protection Act, 2023 (Act No. 22 of 2023). The Gazette of India, Extraordinary, Part II—Section 1."),
        ("World Health Organization. (2021). Ethics and governance of artificial intelligence for health: WHO guidance. World Health Organization. ISBN: 978-92-4-002920-0."),
        ("Indian Council of Medical Research. (2023). Ethical guidelines for application of artificial intelligence in biomedical research and healthcare. ICMR, New Delhi, India."),
        ("OWASP Foundation. (2021). OWASP Top 10:2021 The Ten Most Critical Web Application Security Risks. https://owasp.org/Top10/"),
        ("Dworkin, M. (2007). Recommendation for block cipher modes of operation: Galois/Counter Mode (GCM) and GMAC. NIST Special Publication 800-38D, National Institute of Standards and Technology.")
    ]
    
    for i, full_text in enumerate(refs):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.space_before = Pt(2); p.paragraph_format.space_after = Pt(4); p.paragraph_format.line_spacing = 1.15
        r1 = p.add_run(f"[{i+1}] ")
        r1.font.name = 'Times New Roman'; r1.font.size = Pt(9.5); r1.bold = True; r1.font.color.rgb = NAVY_BLUE
        r2 = p.add_run(full_text)
        r2.font.name = 'Times New Roman'; r2.font.size = Pt(9.5)

def add_appendix_a(doc):
    add_h1(doc, "Appendix A: Engineering Challenges, Root Causes, and Architectural Fixes")
    add_p(doc, "Table A.1 documents ten primary engineering challenges diagnosed and remediated during project development.")

    add_table_caption(doc, "Table A.1: Engineering Challenges, Root Causes, and Architectural Remediation Fixes")
    t_a1 = doc.add_table(rows=len(APPENDIX_A_CHALLENGES)+1, cols=4)
    t_a1.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t_a1)
    cwa1 = [Inches(1.6), Inches(2.3), Inches(2.3), Inches(1.0)]
    
    ha1 = ["Challenge", "Diagnosed Root Cause", "Architectural Remediation Implemented", "Evidence"]
    for j, h in enumerate(ha1):
        cell = t_a1.cell(0, j); cell.width = cwa1[j]
        set_cell_shading(cell, "154360"); set_cell_margins(cell, top=70, bottom=70, left=50, right=50)
        p = cell.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(h); r.font.name = 'Times New Roman'; r.font.size = Pt(8.5); r.bold = True; r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        
    for i, row in enumerate(APPENDIX_A_CHALLENGES):
        row_cells = t_a1.rows[i+1].cells
        for j, w in enumerate(cwa1):
            row_cells[j].width = w; set_cell_margins(row_cells[j], top=50, bottom=50, left=50, right=50)
            if i % 2 == 1: set_cell_shading(row_cells[j], "F4F6F7")
        
        p0 = row_cells[0].paragraphs[0]; r = p0.add_run(row["problem"]); r.font.name = 'Times New Roman'; r.font.size = Pt(8); r.bold = True
        p1 = row_cells[1].paragraphs[0]; r = p1.add_run(row["found"]); r.font.name = 'Times New Roman'; r.font.size = Pt(7.5)
        p2 = row_cells[2].paragraphs[0]; r = p2.add_run(row["fix"]); r.font.name = 'Times New Roman'; r.font.size = Pt(7.5)
        p3 = row_cells[3].paragraphs[0]; r = p3.add_run(row["evidence"]); r.font.name = 'Times New Roman'; r.font.size = Pt(7); r.font.color.rgb = MUTED_GREY

def add_appendix_b(doc):
    add_h1(doc, "Appendix B: Reproduction Commands Playbook")
    add_p(doc, "Listing B.1 provides the sequence of 10 command-line instructions to launch the cluster, execute all 322 automated test suites, run golden test cases, run Playwright journeys, execute load testing, and run security scans.")

    cmds = (
        "# 1. Clone repository and start Docker cluster\n"
        "docker compose -f docker-compose.yml up -d\n"
        "powershell -File scripts/health_check.ps1\n\n"
        "# 2. Seed clean test data\n"
        "powershell -File scripts/seed_docker.ps1\n\n"
        "# 3. Execute Backend API test suite (231 tests)\n"
        "docker compose exec -T backend npm test -- --coverage\n\n"
        "# 4. Execute ML Engine test suite (55 tests)\n"
        "docker compose exec -T ml pytest -v\n\n"
        "# 5. Execute Guarded LLM Graph test suite (26 tests)\n"
        "docker compose exec -T graph pytest -v\n\n"
        "# 6. Execute Frontend UI test suite (10 tests)\n"
        "docker compose exec -T frontend pytest -v\n\n"
        "# 7. Run Golden Clinical Test Set (20 cases)\n"
        "python scripts/test_p0_golden_set.py\n\n"
        "# 8. Run Playwright End-to-End Suite (11 journeys)\n"
        "node scripts/test_p4_playwright_suite.mjs\n\n"
        "# 9. Execute Load Test Suite in test mode (2,400 requests)\n"
        "docker compose -f docker-compose.test.yml up -d\n"
        "python scripts/run_p5_load_test_testmode.py\n\n"
        "# 10. Run Container Security Scan with Trivy\n"
        "powershell -File scripts/run_trivy_scan_all.ps1"
    )

    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_shading(cell, "263238")
    set_cell_margins(cell, top=80, bottom=80, left=100, right=100)
    p = cell.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.LEFT; p.paragraph_format.line_spacing = 1.05
    r = p.add_run(cmds); r.font.name = 'Consolas'; r.font.size = Pt(8); r.font.color.rgb = RGBColor(0xEE, 0xFF, 0x41)
