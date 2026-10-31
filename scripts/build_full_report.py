"""
ArogyaMitra — Complete B.Tech Project Report Builder & PDF Converter
Assembles the complete 50-70 page academic project report (.docx) and converts it to PDF.
Strict adherence to IIIT Pune academic standards, checkable numbers, and literature grounding.
"""

import os
import sys
import docx
import win32com.client
from pdfminer.high_level import extract_pages

# Ensure scripts dir is on sys.path
SCRIPTS_DIR = os.path.dirname(os.path.abspath(__file__))
if SCRIPTS_DIR not in sys.path:
    sys.path.insert(0, SCRIPTS_DIR)

from report_styles import setup_document_geometry
from report_frontmatter import (
    add_cover_page,
    add_bonafide_certificate,
    add_undertaking,
    add_conflict_of_interest,
    add_acknowledgements,
    add_abstract,
    add_table_of_contents_pages,
    add_list_of_figures_and_tables
)
from report_chapters_1_2 import add_chapter_1, add_chapter_2
from report_chapter_3 import add_chapter_3
from report_chapter_4 import add_chapter_4
from report_chapter_5_and_backmatter import (
    add_chapter_5,
    add_references,
    add_appendix_a,
    add_appendix_b,
    add_appendix_c
)

BASE_DIR = os.path.dirname(SCRIPTS_DIR)
OUTPUT_DIR = os.path.join(BASE_DIR, "docs", "REPORT_FINAL")
os.makedirs(OUTPUT_DIR, exist_ok=True)

DOCX_PATH = os.path.join(OUTPUT_DIR, "ArogyaMitra_BTech_Project_Report.docx")
PDF_PATH = os.path.join(OUTPUT_DIR, "ArogyaMitra_BTech_Project_Report.pdf")

def build_report():
    print("=" * 70)
    print("Building ArogyaMitra B.Tech Project Report (.docx)...")
    print("=" * 70)
    
    doc = docx.Document()
    setup_document_geometry(doc)
    
    print("1. Assembling Front Matter (Cover, Certificate, Undertaking, Abstract, TOC)...")
    add_cover_page(doc)
    add_bonafide_certificate(doc)
    add_undertaking(doc)
    add_conflict_of_interest(doc)
    add_acknowledgements(doc)
    add_abstract(doc)
    add_table_of_contents_pages(doc)
    add_list_of_figures_and_tables(doc)
    
    print("2. Assembling Chapter 1: Introduction & Literature Review...")
    add_chapter_1(doc)
    
    print("3. Assembling Chapter 2: System Requirements & Threat Modeling...")
    add_chapter_2(doc)
    
    print("4. Assembling Chapter 3: System Architecture & Design...")
    add_chapter_3(doc)
    
    print("5. Assembling Chapter 4: Empirical Evaluation & Experimental Results...")
    add_chapter_4(doc)
    
    print("6. Assembling Chapter 5: Conclusion & Future Scope...")
    add_chapter_5(doc)
    
    print("7. Assembling Formal References...")
    add_references(doc)
    
    print("8. Assembling Appendix A: Engineering Challenges & Fixes...")
    add_appendix_a(doc)
    
    print("9. Assembling Appendix B: Reproduction Commands Playbook...")
    add_appendix_b(doc)
    
    print("10. Assembling Appendix C: Complete Test Suite Inventory...")
    add_appendix_c(doc)
    
    print(f"Saving Word Document to: {DOCX_PATH}")
    doc.save(DOCX_PATH)
    docx_size = os.path.getsize(DOCX_PATH)
    print(f"Word document saved successfully! Size: {docx_size:,} bytes")
    
    print("=" * 70)
    print("Converting DOCX to PDF using Microsoft Word COM Automation...")
    print("=" * 70)
    
    word = None
    try:
        word = win32com.client.Dispatch("Word.Application")
        word.Visible = False
        word.DisplayAlerts = False
        
        abs_docx = os.path.abspath(DOCX_PATH)
        abs_pdf = os.path.abspath(PDF_PATH)
        
        print(f"Opening: {abs_docx}")
        word_doc = word.Documents.Open(abs_docx)
        
        # wdStatisticPages = 2
        word_pages = word_doc.ComputeStatistics(2)
        print(f"Word Document Computed Statistics - Page Count: {word_pages}")
        
        # FileFormat=17 is wdFormatPDF
        print(f"Exporting to PDF: {abs_pdf}")
        word_doc.SaveAs(abs_pdf, FileFormat=17)
        word_doc.Close()
        
        pdf_size = os.path.getsize(PDF_PATH)
        print(f"PDF exported successfully! Size: {pdf_size:,} bytes")
        
        # Verify page count in generated PDF
        pdf_pages = len(list(extract_pages(PDF_PATH)))
        print(f"PDF Verified Page Count: {pdf_pages} pages")
        
    except Exception as e:
        print(f"Error during Word COM conversion: {e}")
        raise e
    finally:
        if word:
            word.Quit()
            
    print("=" * 70)
    print("ArogyaMitra Project Report Generation Complete!")
    print(f"DOCX: {DOCX_PATH} ({os.path.getsize(DOCX_PATH):,} bytes)")
    print(f"PDF:  {PDF_PATH} ({os.path.getsize(PDF_PATH):,} bytes, {pdf_pages} pages)")
    print("=" * 70)

if __name__ == "__main__":
    build_report()
